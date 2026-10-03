from io import BytesIO

import pymupdf
import pytest
from docx import Document as DocxDocument
from PIL import Image
from pypdf import PdfReader, PdfWriter

from superai_review.documents.errors import PasswordRequiredError
from superai_review.documents.extractor import DocumentExtractor
from superai_review.documents.types import DocumentMime


class FakeOCR:
    def __init__(self) -> None:
        self.image_calls = 0
        self.pdf_calls: list[int] = []

    def image_to_text(self, data: bytes) -> str:
        assert data
        self.image_calls += 1
        return "image text from OCR"

    def pdf_page_to_text(
        self,
        data: bytes,
        page_index: int,
        password: str | None = None,
    ) -> str:
        assert data
        del password
        self.pdf_calls.append(page_index)
        return f"scanned page {page_index + 1}"


def pdf_with_text() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "This PDF contains enough machine-readable text to avoid the OCR fallback.",
    )
    data = document.tobytes()
    document.close()
    return data


def blank_pdf() -> bytes:
    document = pymupdf.open()
    document.new_page()
    data = document.tobytes()
    document.close()
    return data


def encrypted_pdf(password: str) -> bytes:
    reader = PdfReader(BytesIO(pdf_with_text()))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(password)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_text_pdf_does_not_call_ocr() -> None:
    ocr = FakeOCR()
    result = DocumentExtractor(ocr).extract(pdf_with_text(), DocumentMime.PDF)

    assert result.page_count == 1
    assert result.ocr_used is False
    assert result.segments[0].kind == "pdf-text"
    assert ocr.pdf_calls == []


def test_scanned_pdf_calls_ocr_per_page() -> None:
    ocr = FakeOCR()
    result = DocumentExtractor(ocr).extract(blank_pdf(), DocumentMime.PDF)

    assert result.ocr_used is True
    assert result.segments[0].kind == "pdf-ocr"
    assert result.segments[0].text == "scanned page 1"
    assert ocr.pdf_calls == [0]


def test_encrypted_pdf_requires_password_but_does_not_persist_it() -> None:
    ocr = FakeOCR()
    data = encrypted_pdf("one-time-secret")

    with pytest.raises(PasswordRequiredError):
        DocumentExtractor(ocr).extract(data, DocumentMime.PDF)

    result = DocumentExtractor(ocr).extract(data, DocumentMime.PDF, "one-time-secret")
    assert result.segments[0].kind == "pdf-text"
    assert ocr.pdf_calls == []


def test_docx_preserves_headings_paragraphs_and_tables() -> None:
    document = DocxDocument()
    document.add_heading("Heading", level=1)
    document.add_paragraph("Paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    output = BytesIO()
    document.save(output)

    result = DocumentExtractor(FakeOCR()).extract(output.getvalue(), DocumentMime.DOCX)

    assert [segment.kind for segment in result.segments] == ["heading", "paragraph", "table"]
    assert result.segments[-1].text == "A\tB"


def test_image_uses_ocr() -> None:
    image = Image.new("RGB", (20, 20), "white")
    output = BytesIO()
    image.save(output, format="PNG")
    ocr = FakeOCR()

    result = DocumentExtractor(ocr).extract(output.getvalue(), DocumentMime.PNG)

    assert result.ocr_used is True
    assert result.segments[0].text == "image text from OCR"
    assert ocr.image_calls == 1
