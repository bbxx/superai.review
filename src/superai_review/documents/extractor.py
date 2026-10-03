from io import BytesIO
from typing import Iterator

from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentType
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from pypdf.errors import FileNotDecryptedError, PdfReadError

from superai_review.documents.errors import (
    ExtractionFailedError,
    InvalidPasswordError,
    PasswordRequiredError,
)
from superai_review.documents.ocr import OCRService
from superai_review.documents.types import DocumentMime, ExtractedDocument, ExtractedSegment


def _docx_blocks(document: DocxDocumentType) -> Iterator[Paragraph | Table]:
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


class DocumentExtractor:
    def __init__(self, ocr: OCRService, pdf_text_min_chars_per_page: int = 24) -> None:
        self._ocr = ocr
        self._pdf_text_min_chars_per_page = pdf_text_min_chars_per_page

    def extract(
        self,
        data: bytes,
        content_type: DocumentMime,
        password: str | None = None,
    ) -> ExtractedDocument:
        if content_type is DocumentMime.PDF:
            return self._pdf(data, password)
        if content_type is DocumentMime.DOCX:
            return self._docx(data)
        if content_type in {DocumentMime.TEXT, DocumentMime.MARKDOWN}:
            return self._text(data, content_type)
        if content_type in {DocumentMime.PNG, DocumentMime.JPEG}:
            return self._image(data, content_type)
        raise ExtractionFailedError("unsupported extraction type")

    def _pdf(self, data: bytes, password: str | None) -> ExtractedDocument:
        try:
            reader = PdfReader(BytesIO(data), strict=False)
            if reader.is_encrypted:
                if password is None:
                    raise PasswordRequiredError("PDF password is required")
                if not reader.decrypt(password):
                    raise InvalidPasswordError("invalid PDF password")

            segments: list[ExtractedSegment] = []
            ocr_used = False
            for index, page in enumerate(reader.pages):
                text = (page.extract_text() or "").strip()
                compact = "".join(text.split())
                kind = "pdf-text"
                if len(compact) < self._pdf_text_min_chars_per_page:
                    text = self._ocr.pdf_page_to_text(data, index, password).strip()
                    ocr_used = True
                    kind = "pdf-ocr"
                segments.append(
                    ExtractedSegment(
                        ordinal=index,
                        page=index + 1,
                        kind=kind,
                        text=text,
                    )
                )
            if not segments:
                raise ExtractionFailedError("PDF contains no pages")
            return ExtractedDocument(
                content_type=DocumentMime.PDF,
                page_count=len(reader.pages),
                segments=segments,
                ocr_used=ocr_used,
            )
        except (PasswordRequiredError, InvalidPasswordError):
            raise
        except FileNotDecryptedError as exc:
            raise InvalidPasswordError("invalid PDF password") from exc
        except PdfReadError as exc:
            raise ExtractionFailedError("invalid PDF") from exc

    def _docx(self, data: bytes) -> ExtractedDocument:
        try:
            document = DocxDocument(BytesIO(data))
        except Exception as exc:
            raise ExtractionFailedError("invalid DOCX") from exc

        segments: list[ExtractedSegment] = []
        for block in _docx_blocks(document):
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if not text:
                    continue
                style_name = block.style.name if block.style is not None else ""
                kind = "heading" if style_name.casefold().startswith("heading") else "paragraph"
            else:
                rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in block.rows]
                text = "\n".join(row for row in rows if row.strip())
                if not text:
                    continue
                kind = "table"
            segments.append(
                ExtractedSegment(
                    ordinal=len(segments),
                    page=None,
                    kind=kind,
                    text=text,
                )
            )

        return ExtractedDocument(
            content_type=DocumentMime.DOCX,
            page_count=1,
            segments=segments,
        )

    @staticmethod
    def _text(data: bytes, content_type: DocumentMime) -> ExtractedDocument:
        text = data.decode("utf-8")
        return ExtractedDocument(
            content_type=content_type,
            page_count=1,
            segments=[ExtractedSegment(ordinal=0, page=None, kind="text", text=text)],
        )

    def _image(self, data: bytes, content_type: DocumentMime) -> ExtractedDocument:
        text = self._ocr.image_to_text(data).strip()
        return ExtractedDocument(
            content_type=content_type,
            page_count=1,
            segments=[ExtractedSegment(ordinal=0, page=1, kind="image-ocr", text=text)],
            ocr_used=True,
        )
